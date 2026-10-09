import openpyxl
import pandas as pd
import os

print("================================================================================")
print("ICE DATABASE V5.0 AUDIT")
print("================================================================================")

ice_path = 'data/raw/ice_database/ICE DB Educational V5.0 - Jun 2026.xlsx'
print(f"ICE Excel Path: {ice_path}")
print(f"File Size: {os.path.getsize(ice_path)} bytes ({os.path.getsize(ice_path)/(1024*1024):.2f} MB)")

wb = openpyxl.load_workbook(ice_path, data_only=True)
print(f"\nTotal Sheets: {len(wb.sheetnames)}")
print("Sheet Names:")
for i, name in enumerate(wb.sheetnames, 1):
    print(f"  {i}. {name}")

# Read ICE Summary sheet via pandas
df_raw = pd.read_excel(ice_path, sheet_name='ICE Summary', header=None)
print(f"\nICE Summary Raw Shape: {df_raw.shape[0]} rows, {df_raw.shape[1]} columns")

materials = []
current_cat = 'General'

for idx, row in df_raw.iloc[18:].iterrows():
    c5 = row[5]  # Material description or category name
    c6 = row[6]  # Embodied carbon factor (kg CO2e / kg)
    c9 = row[9] if len(row) > 9 else None   # Comments
    c10 = row[10] if len(row) > 10 else None # DQI Score
    c11 = row[11] if len(row) > 11 else None # DQI Version
    c19 = row[19] if len(row) > 19 else None # Full Name in ICE DB
    
    if pd.notnull(c5) and (pd.isnull(c6) or str(c6).strip() == ''):
        s = str(c5).strip()
        if s not in ['Materials', 'Note on DQI Scores', 'Introduction']:
            current_cat = s
    elif pd.notnull(c5) and isinstance(c6, (int, float)):
        materials.append({
            'Category': current_cat,
            'Material': str(c5).strip(),
            'Embodied_Carbon_kgCO2e_kg': float(c6),
            'DQI_Score': float(c10) if pd.notnull(c10) and isinstance(c10, (int, float)) else None,
            'DQI_Version': str(c11).strip() if pd.notnull(c11) else '',
            'Comments': str(c9).strip() if pd.notnull(c9) else '',
            'Full_Name': str(c19).strip() if pd.notnull(c19) else str(c5).strip()
        })

df_ice = pd.DataFrame(materials)
print(f"\nTotal Materials Extracted: {len(df_ice)}")
print("\nMaterials Count by Category:")
print(df_ice['Category'].value_counts().to_string())

# Usable materials for walls and roofs
wall_roof_categories = [
    'Clay Bricks', 'Concrete', 'Cement and Mortar', 'Insulation',
    'Timber', 'Plaster', 'Glass', 'Steel', 'Aluminium', 'Asphalt', 'Bitumen', 'Paint'
]

print("\n================================================================================")
print("USABLE MATERIALS FOR WALLS AND ROOFS")
print("================================================================================")

for cat in wall_roof_categories:
    sub = df_ice[df_ice['Category'].str.lower() == cat.lower()]
    if len(sub) == 0:
        # try partial match
        sub = df_ice[df_ice['Category'].str.contains(cat, case=False, na=False)]
    print(f"\n--- Category: {cat} ({len(sub)} materials) ---")
    display_cols = ['Material', 'Embodied_Carbon_kgCO2e_kg', 'DQI_Score']
    print(sub[display_cols].head(8).to_string(index=False))

# Export a clean summary table to CSV for processed data
out_csv = 'data/processed/ice_materials_extracted.csv'
df_ice.to_csv(out_csv, index=False)
print(f"\nSaved structured ICE materials to: {out_csv}")
